# Copyright 2026 Seedian
# License AGPL-3.0 or later (http://www.gnu.org/licenses/agpl).

import zipfile
from datetime import date, datetime
from io import BytesIO
from itertools import pairwise

from openpyxl import load_workbook

from odoo.tests import TransactionCase

from odoo.addons.collection.report.collection_sheet_xlsx import HDR_BLUE, NAVY_DARK
from odoo.addons.collection_receipt_proposal.report import (
    collection_receipt_assignment_xlsx as builder,
)

from .common import ReceiptProposalCommon

XLSX_MODEL = "collection.receipt.assignment.xlsx"
SUMMARY_HEADER_ROW = 11
SUMMARY_TITLES = [
    "Código",
    "Cliente",
    "Circuito",
    "Cobrado",
    "Aplicado",
    "A Cuenta",
    "Deuda Total",
    "Pendiente",
    "Estado Circuito",
    "Nota de Cierre",
]
# `Documento` spans A:B, so openpyxl reads the merged B cell as None.
DOCUMENT_TITLES = [
    "Documento",
    None,
    "Emisión",
    "Vencimiento",
    "Saldo",
    "Aplicado",
    "Comisión",
    None,
    None,
]
TRANSFER_TITLES = [
    "Referencia",
    "Fecha",
    "Nº Comprobante",
    "Cuenta Origen",
    "Cuenta Destino",
    None,
    None,
    "Importe del Valor",
    "Aplicado",
]
CHECK_TITLES = [
    "Referencia",
    "Emisión",
    "Nº Cheque",
    "Banco",
    "Tipo",
    "Pago",
    None,
    "Importe del Valor",
    "Aplicado",
]
# 1-based openpyxl columns on the Clientes grid.
DOCUMENT_APPLIED_COL = 6
VOUCHER_AMOUNT_COL = 8
VOUCHER_APPLIED_COL = 9
CASH_AMOUNT_COL = 4
COLLECTED_COL = 7
APPLIED_COL = 8
ON_ACCOUNT_COL = 9
# Geometry of the one-customer fixture below: one circuit, two documents, a
# transfer block, a cheque block holding a split, one replacement.
CUSTOMER_BAND_ROW = 6
CIRCUIT_BAND_ROW = 7
DOCUMENT_HEADING_ROW = 8
DOCUMENT_HEADER_ROW = 9
DOCUMENT_FIRST_ROW = 10
DOCUMENT_SUBTOTAL_ROW = 12
TRANSFER_HEADING_ROW = 13
TRANSFER_HEADER_ROW = 14
TRANSFER_FIRST_ROW = 15
TRANSFER_SUBTOTAL_ROW = 16
CHECK_HEADING_ROW = 17
CHECK_HEADER_ROW = 18
CHECK_FIRST_ROW = 19
CHECK_SUBTOTAL_ROW = 20
CASH_HEADING_ROW = 21
CASH_HEADER_ROW = 22
CASH_FIRST_ROW = 23
CASH_SUBTOTAL_ROW = 24
BALANCE_ROW = 25
CUSTOMER_TOTAL_ROW = 27
GRAND_TOTAL_ROW = 29
UNASSIGNED_SECTION_ROW = 31
UNASSIGNED_HEADER_ROW = 33
UNASSIGNED_FIRST_ROW = 34
UNASSIGNED_SUBTOTAL_ROW = 35
UNASSIGNED_TOTAL_ROW = 36
RECON_SECTION_ROW = 38
RECON_ASSIGNED_ROW = 39
RECON_COMPUTED_ROW = 42
RECON_SHEET_ROW = 43
RECON_DIFFERENCE_ROW = 44


def _luminance(colour):
    """WCAG relative luminance of an ``#RRGGBB`` string."""
    colour = colour.lstrip("#")
    channels = [int(colour[index : index + 2], 16) / 255 for index in (0, 2, 4)]
    channels = [
        value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4
        for value in channels
    ]
    return 0.2126 * channels[0] + 0.7152 * channels[1] + 0.0722 * channels[2]


def _contrast(first, second):
    """WCAG contrast ratio between two ``#RRGGBB`` strings."""
    values = sorted((_luminance(first), _luminance(second)))
    return (values[1] + 0.05) / (values[0] + 0.05)


def _summary_row(kind, kind_label, collected, applied, debt, pending, status):
    return {
        "code": "C001",
        "name": "ACME SA",
        "kind": kind,
        "kind_label": kind_label,
        "collected": collected,
        "applied": applied,
        "on_account": collected - applied,
        "debt_total": debt,
        "pending": pending,
        "kind_status_label": status,
        "closed": False,
        "note": "",
    }


def _document_row(number, external, applied):
    return {
        "number": number,
        "external_id": external,
        "issue_date": date(2026, 1, 1),
        "due_date": date(2026, 1, 10),
        "invoice_amount": 100.0,
        "invoice_balance": 100.0,
        "amount_applied": applied,
        "commission": applied * 0.05,
        "invoice_state_label": "Vigente",
    }


def _value_row(reference, voucher_amount, applied, **extra):
    row = {
        "reference": reference,
        "document_id": f"DOC-{reference}",
        "partner_name": "ACME SA",
        "type_label": "Transferencia bancaria",
        "date": date(2026, 9, 2),
        "number": "TR-1",
        "voucher_amount": voucher_amount,
        "applied_amount": applied,
        "state_label": "Validado por cobranzas",
        "check_bank": "",
        "check_type_label": "",
        "check_issue_date": False,
        "check_payment_date": False,
        "from_account": "0001",
        "to_account": "0002",
        "retention_label": "",
        "reviewed_by": "Revisor Uno",
        "reviewed_at": datetime(2026, 9, 3, 10, 30),
        "is_partial": voucher_amount != applied,
    }
    row.update(extra)
    return row


def _voucher_block(code, label, rows, **extra):
    block = {
        "code": code,
        "label": label,
        "rows": rows,
        "subtotal_applied": sum(row["applied_amount"] for row in rows),
        "subtotal_voucher_amount": sum(row["voucher_amount"] for row in rows),
        "split_count": sum(1 for row in rows if row["is_partial"]),
    }
    block.update(extra)
    return block


def _fixture_payload():
    """A payload of primitives only — the builder never sees the ORM."""
    rows = [
        _summary_row(
            "invoice", "Facturas", 150.0, 150.0, 200.0, 50.0, "Deuda pendiente"
        ),
        _summary_row("quotation", "Presupuestos", 50.0, 50.0, 50.0, 0.0, "Conciliado"),
    ]
    money_keys = ("collected", "applied", "on_account", "debt_total", "pending")
    circuit = {
        "kind": "invoice",
        "kind_label": "Facturas",
        "receipt_reference": "REC-0007",
        "receipt_state": "draft",
        "receipt_state_label": "Borrador",
        "total_values": 150.0,
        "total_applied": 150.0,
        "amount_on_account": 0.0,
        "document_rows": [
            _document_row("FA-0001", "A-OLD", 100.0),
            _document_row("FA-0002", "A-NEW", 50.0),
        ],
        "subtotal_invoice_balance": 200.0,
        "subtotal_amount_applied": 150.0,
        "subtotal_commission": 7.5,
        "voucher_blocks": [
            _voucher_block(
                "bank_transfer",
                "Transferencia bancaria",
                [_value_row("V-0001", 100.0, 100.0)],
            ),
            _voucher_block(
                "check",
                "Cheque",
                [
                    _value_row(
                        "V-0002",
                        300.0,
                        50.0,
                        type_label="Cheque",
                        number="CH-77",
                        check_bank="Banco Nación",
                        check_type_label="Diferido",
                        check_issue_date=date(2026, 9, 1),
                        check_payment_date=date(2026, 10, 1),
                    )
                ],
            ),
        ],
        "subtotal_values_applied": 150.0,
        "split_count": 1,
        "cash_rows": [
            {
                "cash_reference": "V-CASH",
                "deposit_reference": "V-DEP",
                "amount": 25.0,
                "direction": "replace",
                "direction_label": "Reemplazo",
                "created_by": "Operador Uno",
                "created_at": datetime(2026, 9, 4, 9, 15),
            }
        ],
        "subtotal_cash_amount": 25.0,
    }
    customer = {
        "code": "C001",
        "name": "ACME SA",
        "closed": False,
        "note": "",
        "circuits": [circuit],
        "total_collected": 150.0,
        "total_applied": 150.0,
        "total_on_account": 0.0,
        "total_commission": 7.5,
    }
    return {
        "company": "ACME",
        "printed_on": date(2026, 9, 8),
        "reference": "CS/2026/0001",
        "state_label": "Propuesta",
        "seller_name": "Vendedor Uno",
        "currency_name": "ARS",
        "date_from": date(2026, 9, 1),
        "date_to": date(2026, 9, 8),
        "customer_count": 1,
        "commission_percent": 5.0,
        "summary_rows": rows,
        "summary_totals": {key: sum(row[key] for row in rows) for key in money_keys},
        "kind_totals": {
            row["kind"]: {key: row[key] for key in money_keys} for row in rows
        },
        "customer_blocks": [customer],
        "customer_totals": {
            "total_collected": 150.0,
            "total_applied": 150.0,
            "total_on_account": 0.0,
        },
        "unassigned": {
            "blocks": [
                _voucher_block(
                    "bank_transfer",
                    "Transferencia bancaria",
                    [_value_row("V-0003", 300.0, 250.0)],
                )
            ],
            "subtotal_voucher_amount": 300.0,
            "subtotal_free_amount": 250.0,
        },
        "recon": {
            "assigned_total": 150.0,
            "unassigned_total": 250.0,
            "replaced_total": 25.0,
            "computed_total": 425.0,
            "sheet_total": 425.0,
            "difference": 0.0,
            "status": "balanced",
        },
    }


class TestReceiptAssignmentXlsxBuilder(TransactionCase):
    """RED/GREEN: the workbook layout, driven by a hand-made payload."""

    def setUp(self):
        super().setUp()
        self.payload = _fixture_payload()
        self.customer = self.payload["customer_blocks"][0]
        self.circuit = self.customer["circuits"][0]
        self.content = builder._build_xlsx(self.payload)

    def _load(self, data_only=False):
        return load_workbook(BytesIO(self.content), data_only=data_only)

    def _titles(self, worksheet, row):
        return [cell.value for cell in worksheet[row]]

    def _values(self, worksheet, row):
        """openpyxl gives ints for whole numbers and None for blank cells."""
        return [
            float(cell.value)
            if isinstance(cell.value, (int, float))
            else cell.value or ""
            for cell in worksheet[row]
        ]

    # ------------------------------------------------------------------
    # Workbook and Resumen
    # ------------------------------------------------------------------

    def test_bytes_are_a_zip_with_resumen_and_clientes(self):
        self.assertTrue(self.content.startswith(b"PK"), self.content[:8])
        with zipfile.ZipFile(BytesIO(self.content)) as zipped:
            self.assertIn("xl/workbook.xml", zipped.namelist())
        self.assertEqual(self._load().sheetnames, ["Resumen", "Clientes"])

    def test_summary_headers_and_first_data_row(self):
        worksheet = self._load()["Resumen"]
        self.assertEqual(self._titles(worksheet, SUMMARY_HEADER_ROW), SUMMARY_TITLES)
        self.assertEqual(
            self._values(worksheet, SUMMARY_HEADER_ROW + 1),
            [
                "C001",
                "ACME SA",
                "Facturas",
                150.0,
                150.0,
                0.0,
                200.0,
                50.0,
                "Deuda pendiente",
                "",
            ],
        )

    def test_summary_footer_has_kind_subtotals_and_a_total(self):
        worksheet = self._load()["Resumen"]
        cached = self._load(data_only=True)["Resumen"]
        first_footer = SUMMARY_HEADER_ROW + 1 + len(self.payload["summary_rows"])
        self.assertEqual(worksheet.cell(first_footer, 1).value, "Subtotal Facturas")
        self.assertEqual(
            worksheet.cell(first_footer + 1, 1).value, "Subtotal Presupuestos"
        )
        self.assertEqual(worksheet.cell(first_footer + 2, 1).value, "Total")
        self.assertIn("SUMIF", worksheet.cell(first_footer, 4).value.upper())
        self.assertIn("SUM(", worksheet.cell(first_footer + 2, 4).value.upper())
        self.assertEqual(
            cached.cell(first_footer, 4).value,
            self.payload["kind_totals"]["invoice"]["collected"],
        )
        self.assertEqual(
            cached.cell(first_footer + 2, 4).value,
            self.payload["summary_totals"]["collected"],
        )

    # ------------------------------------------------------------------
    # Clientes: the three heading levels
    # ------------------------------------------------------------------

    def test_customer_band_names_the_customer_without_a_status(self):
        worksheet = self._load()["Clientes"]
        self.assertEqual(
            worksheet.cell(CUSTOMER_BAND_ROW, 1).value,
            "Cliente — C001  ACME SA",
        )

    def test_circuit_band_names_the_receipt_and_its_state(self):
        worksheet = self._load()["Clientes"]
        self.assertEqual(
            worksheet.cell(CIRCUIT_BAND_ROW, 1).value,
            "Facturas — Recibo REC-0007 (Borrador)",
        )

    def _fill(self, worksheet, row, col=1):
        rgb = worksheet.cell(row, col).fill.start_color.rgb
        return (rgb or "")[-6:] if isinstance(rgb, str) else ""

    def test_the_four_heading_levels_each_have_their_own_fill(self):
        """RED/GREEN: customer, circuit, block and header must never collide."""
        worksheet = self._load()["Clientes"]
        fills = {
            "customer": self._fill(worksheet, CUSTOMER_BAND_ROW),
            "circuit": self._fill(worksheet, CIRCUIT_BAND_ROW),
            "block": self._fill(worksheet, DOCUMENT_HEADING_ROW),
            "header": self._fill(worksheet, DOCUMENT_HEADER_ROW),
        }
        self.assertEqual(len(set(fills.values())), 4, fills)
        # ... and they must lighten monotonically as the nesting deepens.
        ladder = [
            fills["customer"],
            fills["circuit"],
            fills["block"],
            fills["header"],
        ]
        levels = [_luminance(f"#{value}") for value in ladder]
        self.assertEqual(levels, sorted(levels), ladder)

    def test_palette_separates_every_level_and_clears_wcag_aa(self):
        """RED/GREEN: the ladder is the design; pin it against silent drift."""
        ladder = [
            builder.LVL1_FILL,
            builder.LVL2_FILL,
            builder.LVL3_FILL,
            HDR_BLUE,
            "#FFFFFF",
        ]
        levels = [_luminance(colour) for colour in ladder]
        self.assertEqual(levels, sorted(levels), ladder)
        for darker, lighter in pairwise(ladder):
            self.assertGreaterEqual(_contrast(darker, lighter), 1.25, (darker, lighter))
        for fill, text in (
            (builder.LVL1_FILL, "#FFFFFF"),
            (builder.LVL2_FILL, "#FFFFFF"),
            (builder.LVL3_FILL, NAVY_DARK),
            (builder.BALANCE_FILL, "#FFFFFF"),
            (builder.SUB_FILL, NAVY_DARK),
            (HDR_BLUE, NAVY_DARK),
            (builder.FLAG_FILL, builder.FLAG_FONT),
        ):
            self.assertGreaterEqual(_contrast(fill, text), 4.5, (fill, text))

    def test_zebra_stripe_is_a_hint_not_a_band(self):
        """RED/GREEN: subtle enough not to compete with any heading."""
        self.assertLess(_contrast("#FFFFFF", builder.ZEBRA_FILL), 1.2)
        self.assertGreater(_contrast("#FFFFFF", builder.ZEBRA_FILL), 1.03)
        worksheet = self._load()["Clientes"]
        self.assertNotEqual(
            self._fill(worksheet, DOCUMENT_FIRST_ROW),
            self._fill(worksheet, DOCUMENT_FIRST_ROW + 1),
        )

    def test_totals_escalate_from_subtotal_to_grand_total(self):
        """RED/GREEN: four total weights, each readable as its own rank."""
        worksheet = self._load()["Clientes"]
        subtotal = worksheet.cell(DOCUMENT_SUBTOTAL_ROW, 1)
        balance = worksheet.cell(BALANCE_ROW, 1)
        customer = worksheet.cell(CUSTOMER_TOTAL_ROW, 1)
        grand = worksheet.cell(GRAND_TOTAL_ROW, 1)
        self.assertEqual(
            len(
                {
                    self._fill(worksheet, DOCUMENT_SUBTOTAL_ROW),
                    self._fill(worksheet, BALANCE_ROW),
                    self._fill(worksheet, CUSTOMER_TOTAL_ROW),
                }
            ),
            3,
        )
        # A quiet subtotal closed by a rule, not by a fill that shouts.
        self.assertTrue(subtotal.border.top.style, "subtotal needs its top rule")
        self.assertTrue(balance.font.bold)
        # The grand total outranks the customer total by size, sharing its navy.
        self.assertEqual(
            self._fill(worksheet, GRAND_TOTAL_ROW),
            self._fill(worksheet, CUSTOMER_TOTAL_ROW),
        )
        self.assertGreater(grand.font.size, customer.font.size)

    def test_reconciliation_components_stay_quieter_than_their_total(self):
        """RED/GREEN: the three inputs must not merge into the computed rows."""
        worksheet = self._load()["Clientes"]
        self.assertNotEqual(
            self._fill(worksheet, RECON_ASSIGNED_ROW),
            self._fill(worksheet, RECON_COMPUTED_ROW),
        )

    def test_difference_row_wears_the_banner_colour(self):
        """RED/GREEN: a sheet that does not close says so in red."""
        worksheet = self._load()["Clientes"]
        self.assertEqual(self._fill(worksheet, RECON_DIFFERENCE_ROW), "C6EFCE")
        payload = _fixture_payload()
        payload["recon"].update(difference=25.0, status="outstanding")
        failing = load_workbook(BytesIO(builder._build_xlsx(payload)))["Clientes"]
        rgb = failing.cell(RECON_DIFFERENCE_ROW, 1).fill.start_color.rgb
        self.assertEqual(rgb[-6:], "FFC7CE")

    def test_a_blank_row_separates_a_circuit_from_what_follows(self):
        """RED/GREEN: two receipts must never abut."""
        worksheet = self._load()["Clientes"]
        self.assertIsNone(worksheet.cell(BALANCE_ROW + 1, 1).value)
        self.assertTrue(worksheet.cell(CUSTOMER_TOTAL_ROW, 1).value.startswith("Total"))

    def test_blank_rows_separate_one_customer_from_the_next(self):
        payload = _fixture_payload()
        second = dict(payload["customer_blocks"][0], code="C002", name="BETA SA")
        payload["customer_blocks"].append(second)
        worksheet = load_workbook(BytesIO(builder._build_xlsx(payload)))["Clientes"]
        next_band = CUSTOMER_TOTAL_ROW + 5
        for row in range(CUSTOMER_TOTAL_ROW + 1, next_band):
            self.assertIsNone(worksheet.cell(row, 1).value)
        self.assertEqual(worksheet.cell(next_band, 1).value, "Cliente — C002  BETA SA")

    # ------------------------------------------------------------------
    # Clientes: documents
    # ------------------------------------------------------------------

    def test_document_block_headers_rows_and_subtotal(self):
        worksheet = self._load()["Clientes"]
        cached = self._load(data_only=True)["Clientes"]
        self.assertEqual(
            worksheet.cell(DOCUMENT_HEADING_ROW, 1).value, "Documentos imputados"
        )
        self.assertEqual(self._titles(worksheet, DOCUMENT_HEADER_ROW), DOCUMENT_TITLES)
        self.assertEqual(
            [worksheet.cell(DOCUMENT_FIRST_ROW + offset, 1).value for offset in (0, 1)],
            ["FA-0001", "FA-0002"],
        )
        self.assertEqual(
            worksheet.cell(DOCUMENT_FIRST_ROW, DOCUMENT_APPLIED_COL).value, 100.0
        )
        self.assertEqual(worksheet.cell(DOCUMENT_SUBTOTAL_ROW, 1).value, "Subtotal")
        self.assertIn(
            "SUM(",
            worksheet.cell(DOCUMENT_SUBTOTAL_ROW, DOCUMENT_APPLIED_COL).value.upper(),
        )
        self.assertEqual(
            cached.cell(DOCUMENT_SUBTOTAL_ROW, DOCUMENT_APPLIED_COL).value,
            self.circuit["subtotal_amount_applied"],
        )
        self.assertEqual(
            cached.cell(DOCUMENT_SUBTOTAL_ROW, 7).value,
            self.circuit["subtotal_commission"],
        )

    # ------------------------------------------------------------------
    # Clientes: vouchers, one column shape per type
    # ------------------------------------------------------------------

    def test_transfer_block_shows_the_accounts(self):
        worksheet = self._load()["Clientes"]
        self.assertEqual(
            worksheet.cell(TRANSFER_HEADING_ROW, 1).value,
            "Comprobantes — Transferencia bancaria",
        )
        self.assertEqual(self._titles(worksheet, TRANSFER_HEADER_ROW), TRANSFER_TITLES)
        self.assertEqual(
            self._values(worksheet, TRANSFER_FIRST_ROW),
            [
                "V-0001",
                worksheet.cell(TRANSFER_FIRST_ROW, 2).value,
                "TR-1",
                "0001",
                "0002",
                "",
                "",
                100.0,
                100.0,
            ],
        )

    def test_check_block_reshapes_to_the_cheque_fields(self):
        """RED/GREEN: a cheque prints bank, type and both dates, not accounts."""
        worksheet = self._load()["Clientes"]
        heading = worksheet.cell(CHECK_HEADING_ROW, 1).value
        self.assertTrue(heading.startswith("Comprobantes — Cheque"), heading)
        self.assertIn("1 parcial(es)", heading)
        self.assertEqual(self._titles(worksheet, CHECK_HEADER_ROW), CHECK_TITLES)
        self.assertEqual(worksheet.cell(CHECK_FIRST_ROW, 3).value, "CH-77")
        self.assertEqual(worksheet.cell(CHECK_FIRST_ROW, 4).value, "Banco Nación")
        self.assertEqual(worksheet.cell(CHECK_FIRST_ROW, 5).value, "Diferido")
        self.assertEqual(
            worksheet.cell(CHECK_FIRST_ROW, 2).value.date(), date(2026, 9, 1)
        )
        self.assertEqual(
            worksheet.cell(CHECK_FIRST_ROW, 6).value.date(), date(2026, 10, 1)
        )

    def test_no_block_prints_ids_states_or_reviews(self):
        payload = _fixture_payload()
        payload["customer_blocks"][0]["circuits"][0]["voucher_blocks"] = [
            _voucher_block(code, code, [_value_row(f"V-{code}", 10.0, 10.0)])
            for code in (
                "cash",
                "check",
                "bank_transfer",
                "echeck",
                "retention_certificate",
            )
        ]
        worksheet = load_workbook(BytesIO(builder._build_xlsx(payload)))["Clientes"]
        removed = {
            "ID Externo",
            "Estado ERP",
            "Estado",
            "Revisado Por",
            "Revisado En",
            "ID Documento",
        }
        values = {
            cell.value
            for row in worksheet.iter_rows(min_row=CUSTOMER_BAND_ROW)
            for cell in row
        }
        self.assertFalse(removed & values)
        dropped = {"A-OLD", "Vigente", "DOC-V-cash", "Validado por cobranzas"}
        self.assertFalse(dropped & values)

    def test_split_value_row_is_highlighted(self):
        worksheet = self._load()["Clientes"]
        plain = worksheet.cell(TRANSFER_FIRST_ROW, VOUCHER_APPLIED_COL)
        split = worksheet.cell(CHECK_FIRST_ROW, VOUCHER_APPLIED_COL)
        self.assertNotEqual(split.fill.start_color.rgb, plain.fill.start_color.rgb)
        self.assertTrue(split.font.bold)
        self.assertFalse(plain.font.bold)

    def test_voucher_subtotal_sums_the_applied_column_only(self):
        worksheet = self._load()["Clientes"]
        cached = self._load(data_only=True)["Clientes"]
        for subtotal, block in (
            (TRANSFER_SUBTOTAL_ROW, self.circuit["voucher_blocks"][0]),
            (CHECK_SUBTOTAL_ROW, self.circuit["voucher_blocks"][1]),
        ):
            self.assertEqual(worksheet.cell(subtotal, 1).value, "Subtotal")
            # The label merge swallows `Importe del Valor`, so a split voucher
            # has nowhere to be double-counted.
            self.assertIsNone(worksheet.cell(subtotal, VOUCHER_AMOUNT_COL).value)
            self.assertIn(
                "SUM(", worksheet.cell(subtotal, VOUCHER_APPLIED_COL).value.upper()
            )
            self.assertEqual(
                cached.cell(subtotal, VOUCHER_APPLIED_COL).value,
                block["subtotal_applied"],
            )

    def test_cash_replacement_amount_has_its_own_column(self):
        worksheet = self._load()["Clientes"]
        cached = self._load(data_only=True)["Clientes"]
        self.assertEqual(
            worksheet.cell(CASH_HEADING_ROW, 1).value, "Reemplazos de efectivo"
        )
        self.assertEqual(worksheet.cell(CASH_HEADER_ROW, 1).value, "Valor Efectivo")
        self.assertEqual(worksheet.cell(CASH_HEADER_ROW, 2).value, "Valor Depósito")
        self.assertEqual(worksheet.cell(CASH_FIRST_ROW, 1).value, "V-CASH")
        self.assertEqual(worksheet.cell(CASH_FIRST_ROW, 3).value, "Reemplazo")
        self.assertEqual(worksheet.cell(CASH_FIRST_ROW, CASH_AMOUNT_COL).value, 25.0)
        # Never in a column any balance or reconciliation figure sums.
        self.assertNotIn(CASH_AMOUNT_COL, (DOCUMENT_APPLIED_COL, VOUCHER_APPLIED_COL))
        self.assertEqual(
            cached.cell(CASH_SUBTOTAL_ROW, CASH_AMOUNT_COL).value,
            self.circuit["subtotal_cash_amount"],
        )

    # ------------------------------------------------------------------
    # Clientes: the chain of balances
    # ------------------------------------------------------------------

    def test_circuit_balance_reads_the_subtotals_above_it(self):
        worksheet = self._load()["Clientes"]
        cached = self._load(data_only=True)["Clientes"]
        self.assertEqual(
            worksheet.cell(BALANCE_ROW, 1).value,
            "Balance Facturas — Cobrado / Aplicado / A Cuenta",
        )
        self.assertEqual(
            worksheet.cell(BALANCE_ROW, COLLECTED_COL).value,
            f"=I{TRANSFER_SUBTOTAL_ROW}+I{CHECK_SUBTOTAL_ROW}",
        )
        self.assertEqual(
            worksheet.cell(BALANCE_ROW, APPLIED_COL).value,
            f"=F{DOCUMENT_SUBTOTAL_ROW}",
        )
        self.assertEqual(
            worksheet.cell(BALANCE_ROW, ON_ACCOUNT_COL).value,
            f"=G{BALANCE_ROW}-H{BALANCE_ROW}",
        )
        self.assertEqual(
            cached.cell(BALANCE_ROW, COLLECTED_COL).value, self.circuit["total_values"]
        )
        self.assertEqual(
            cached.cell(BALANCE_ROW, APPLIED_COL).value, self.circuit["total_applied"]
        )
        self.assertEqual(
            cached.cell(BALANCE_ROW, ON_ACCOUNT_COL).value,
            self.circuit["amount_on_account"],
        )

    def test_customer_total_adds_its_circuit_balances(self):
        worksheet = self._load()["Clientes"]
        cached = self._load(data_only=True)["Clientes"]
        self.assertEqual(
            worksheet.cell(CUSTOMER_TOTAL_ROW, 1).value,
            "Total C001  ACME SA — Cobrado / Aplicado / A Cuenta",
        )
        self.assertEqual(
            worksheet.cell(CUSTOMER_TOTAL_ROW, COLLECTED_COL).value,
            f"=G{BALANCE_ROW}",
        )
        self.assertEqual(
            cached.cell(CUSTOMER_TOTAL_ROW, COLLECTED_COL).value,
            self.customer["total_collected"],
        )

    def test_grand_total_adds_the_customer_totals(self):
        worksheet = self._load()["Clientes"]
        cached = self._load(data_only=True)["Clientes"]
        self.assertEqual(
            worksheet.cell(GRAND_TOTAL_ROW, 1).value,
            "Total General — Cobrado / Aplicado / A Cuenta",
        )
        self.assertEqual(
            worksheet.cell(GRAND_TOTAL_ROW, COLLECTED_COL).value,
            f"=G{CUSTOMER_TOTAL_ROW}",
        )
        self.assertEqual(
            cached.cell(GRAND_TOTAL_ROW, APPLIED_COL).value,
            self.payload["customer_totals"]["total_applied"],
        )

    # ------------------------------------------------------------------
    # Clientes: unassigned money and reconciliation
    # ------------------------------------------------------------------

    def test_unassigned_section_reuses_the_type_shapes_and_reads_libre(self):
        worksheet = self._load()["Clientes"]
        cached = self._load(data_only=True)["Clientes"]
        self.assertEqual(
            worksheet.cell(UNASSIGNED_SECTION_ROW, 1).value, "Valores sin asignar"
        )
        titles = self._titles(worksheet, UNASSIGNED_HEADER_ROW)
        self.assertEqual(titles[:7], TRANSFER_TITLES[:7])
        self.assertEqual(titles[VOUCHER_APPLIED_COL - 1], "Libre")
        self.assertEqual(worksheet.cell(UNASSIGNED_FIRST_ROW, 1).value, "V-0003")
        self.assertEqual(
            worksheet.cell(UNASSIGNED_TOTAL_ROW, 1).value, "Total sin asignar"
        )
        self.assertEqual(
            worksheet.cell(UNASSIGNED_TOTAL_ROW, VOUCHER_APPLIED_COL).value,
            f"=I{UNASSIGNED_SUBTOTAL_ROW}",
        )
        self.assertEqual(
            cached.cell(UNASSIGNED_TOTAL_ROW, VOUCHER_APPLIED_COL).value,
            self.payload["unassigned"]["subtotal_free_amount"],
        )

    def test_recon_block_closes_against_the_sheet_total(self):
        worksheet = self._load()["Clientes"]
        cached = self._load(data_only=True)["Clientes"]
        recon = self.payload["recon"]
        self.assertEqual(
            worksheet.cell(RECON_SECTION_ROW, 1).value, "Conciliación con la planilla"
        )
        self.assertEqual(
            [
                worksheet.cell(RECON_ASSIGNED_ROW + offset, 1).value
                for offset in range(6)
            ],
            [
                "Asignado a recibos",
                "Sin asignar",
                "Efectivo reemplazado",
                "Total calculado",
                "Total de la planilla",
                "Diferencia",
            ],
        )
        col = VOUCHER_APPLIED_COL
        self.assertEqual(
            worksheet.cell(RECON_ASSIGNED_ROW, col).value, f"=G{GRAND_TOTAL_ROW}"
        )
        self.assertEqual(
            worksheet.cell(RECON_ASSIGNED_ROW + 1, col).value,
            f"=I{UNASSIGNED_TOTAL_ROW}",
        )
        self.assertEqual(
            worksheet.cell(RECON_ASSIGNED_ROW + 2, col).value, recon["replaced_total"]
        )
        self.assertEqual(
            worksheet.cell(RECON_COMPUTED_ROW, col).value,
            f"=I{RECON_ASSIGNED_ROW}+I{RECON_ASSIGNED_ROW + 1}"
            f"+I{RECON_ASSIGNED_ROW + 2}",
        )
        self.assertEqual(
            worksheet.cell(RECON_DIFFERENCE_ROW, col).value,
            f"=I{RECON_COMPUTED_ROW}-I{RECON_SHEET_ROW}",
        )
        self.assertEqual(
            cached.cell(RECON_COMPUTED_ROW, col).value, recon["computed_total"]
        )
        self.assertEqual(
            cached.cell(RECON_DIFFERENCE_ROW, col).value, recon["difference"]
        )

    # ------------------------------------------------------------------
    # Formats, page setup and geometry
    # ------------------------------------------------------------------

    def test_dates_are_real_dates_and_money_is_formatted(self):
        worksheet = self._load()["Clientes"]
        due = worksheet.cell(DOCUMENT_FIRST_ROW, 4)
        self.assertEqual(due.value.date(), date(2026, 1, 10))
        self.assertEqual(due.number_format, "DD/MM/YYYY")
        self.assertEqual(
            worksheet.cell(DOCUMENT_FIRST_ROW, DOCUMENT_APPLIED_COL).number_format,
            "#,##0.00",
        )
        reviewed = worksheet.cell(CASH_FIRST_ROW, 7)
        self.assertEqual(
            reviewed.value,
            datetime(2026, 9, 4, 9, 15),
        )
        self.assertEqual(reviewed.number_format, "DD/MM/YYYY HH:MM")

    def test_clientes_has_no_autofilter_but_resumen_does(self):
        """RED/GREEN: filtering would hide rows the SUMs still count."""
        workbook = self._load()
        self.assertIsNone(workbook["Clientes"].auto_filter.ref)
        self.assertTrue(workbook["Resumen"].auto_filter.ref)

    def test_page_setup_freezes_each_sheet_where_it_needs_it(self):
        workbook = self._load()
        summary = workbook["Resumen"]
        customers = workbook["Clientes"]
        self.assertEqual(summary.freeze_panes, f"A{SUMMARY_HEADER_ROW + 1}")
        self.assertEqual(summary.page_setup.orientation, "landscape")
        self.assertEqual(summary.page_setup.paperSize, 9)
        self.assertTrue(summary.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(summary.page_setup.fitToHeight, 0)
        self.assertEqual(
            (summary.print_title_rows or "").replace("$", ""),
            f"{SUMMARY_HEADER_ROW}:{SUMMARY_HEADER_ROW}",
        )
        # Clientes has no single header row, so it freezes and repeats the
        # membrete instead.
        self.assertEqual(customers.freeze_panes, "A6")
        self.assertEqual((customers.print_title_rows or "").replace("$", ""), "1:5")
        # xlsxwriter pads the stored width; only the Planilla rewrites its XML
        # to byte-exact widths, and this report has no template to match.
        self.assertEqual(int(summary.column_dimensions["B"].width), 32)

    def test_membrete_carries_the_sheet_context(self):
        worksheet = self._load()["Resumen"]
        # Like the Planilla since `collection` dropped its company header and
        # print date: the membrete is the report name and the sheet reference.
        self.assertEqual(worksheet["A1"].value, "Asignaciones de Cobranza")
        self.assertEqual(worksheet["A2"].value, "CS/2026/0001")
        self.assertFalse(worksheet["A4"].value)
        self.assertEqual(worksheet["A7"].value, "Planilla")
        self.assertEqual(worksheet["C7"].value, "CS/2026/0001")
        self.assertEqual(worksheet["E7"].value, "Estado")
        self.assertEqual(worksheet["G7"].value, "Propuesta")
        self.assertEqual(worksheet["I7"].value, "Moneda")
        self.assertEqual(worksheet["J7"].value, "ARS")
        self.assertEqual(worksheet["A8"].value, "Vendedor")
        self.assertEqual(worksheet["C8"].value, "Vendedor Uno")
        self.assertEqual(worksheet["E8"].value, "Período")
        self.assertEqual(worksheet["G8"].value, "01/09/2026 - 08/09/2026")
        self.assertEqual(worksheet["I8"].value, "Clientes")
        self.assertEqual(worksheet["J8"].value, 1)
        customers = self._load()["Clientes"]
        self.assertEqual(customers["A1"].value, "Asignaciones - Clientes")
        self.assertEqual(customers["A2"].value, "CS/2026/0001")
        self.assertFalse(customers["A4"].value)

    def test_membrete_has_no_company_nor_printed_date(self):
        """RED/GREEN: the header matches the Planilla's, which shows neither."""
        workbook = self._load()
        printed = self.payload["printed_on"].strftime("%d/%m/%Y")
        # Each sheet spans its membrete over its own width: ten columns on
        # Resumen, nine on Clientes.
        for name, last in (("Resumen", "J"), ("Clientes", "I")):
            worksheet = workbook[name]
            merges = {str(cell_range) for cell_range in worksheet.merged_cells.ranges}
            for row in (1, 2, 3, 4):
                self.assertIn(f"A{row}:{last}{row}", merges, name)
            values = [
                cell.value
                for row in worksheet.iter_rows(max_row=5)
                for cell in row
                if cell.value not in (None, "")
            ]
            self.assertNotIn(self.payload["company"], values, name)
            self.assertNotIn(printed, values, name)

    def test_geometry_stays_inside_the_shared_a1_helper(self):
        """RED/GREEN: `_a1` is letter-based, so no formula may pass column Z."""
        self.assertLessEqual(builder.CUSTOMER_LAST_COL, 25)
        self.assertEqual(len(builder.CUSTOMER_WIDTHS), builder.CUSTOMER_LAST_COL + 1)
        self.assertEqual(len(builder.SUMMARY_WIDTHS), builder.SUMMARY_LAST_COL + 1)

    def test_every_voucher_shape_fits_the_grid(self):
        """RED/GREEN: a shape must fit the middle columns, never spill past them."""
        shapes = [builder._DEFAULT_VOUCHER_SHAPE, *builder._VOUCHER_SHAPES.values()]
        middle_slots = builder._VOUCHER_AMOUNT_COL - 2
        for shape in shapes:
            self.assertLessEqual(len(shape["middle"]), middle_slots, shape)
            self.assertEqual(len(shape["date"]), 2, shape)
            for _title, _key, kind in shape["middle"]:
                self.assertIn(kind, ("text", "date", "datetime"), shape)

    def test_empty_payload_still_renders(self):
        """RED/GREEN: a sheet with nothing assigned MUST NOT crash the export."""
        content = builder._build_xlsx({})
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.sheetnames, ["Resumen", "Clientes"])
        summary = workbook["Resumen"]
        self.assertEqual(
            summary.cell(SUMMARY_HEADER_ROW + 1, 1).value, "Subtotal Facturas"
        )
        self.assertEqual(summary.cell(SUMMARY_HEADER_ROW + 3, 1).value, "Total")
        self.assertEqual(summary.cell(SUMMARY_HEADER_ROW + 3, 4).value, 0)
        customers = workbook["Clientes"]
        self.assertEqual(
            customers.cell(6, 1).value, "Total General — Cobrado / Aplicado / A Cuenta"
        )
        self.assertEqual(customers.cell(6, COLLECTED_COL).value, 0)
        self.assertEqual(customers.cell(9, 1).value, "Total sin asignar")
        self.assertEqual(customers.cell(9, VOUCHER_APPLIED_COL).value, 0)
        self.assertEqual(customers.cell(17, 1).value, "Diferencia")


class TestReceiptAssignmentXlsxRender(ReceiptProposalCommon, TransactionCase):
    """RED/GREEN: the AbstractModel renders the sheet's own payload."""

    def setUp(self):
        super().setUp()
        self.scenario = self._board_scenario()
        self.sheet = self.scenario["sheet"]
        allocation = self.scenario["allocation_a"]
        receipt = self._create_receipt(allocation)
        self._create_value(receipt, self.scenario["voucher_a"], 150.0)
        allocation._proposal_recompute("invoice")

    def test_render_xlsx_returns_a_workbook_for_the_live_payload(self):
        payload = self.sheet._receipt_assignment_report_data()
        content = self.env[XLSX_MODEL]._render_xlsx(payload)
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.sheetnames, ["Resumen", "Clientes"])
        summary = workbook["Resumen"]
        self.assertEqual(summary.cell(SUMMARY_HEADER_ROW, 1).value, SUMMARY_TITLES[0])
        names = {
            summary.cell(SUMMARY_HEADER_ROW + 1 + index, 2).value
            for index in range(len(payload["summary_rows"]))
        }
        self.assertIn(self.scenario["allocation_a"].partner_id.display_name, names)
        customers = workbook["Clientes"]
        band = customers.cell(CUSTOMER_BAND_ROW, 1).value
        self.assertIn(self.scenario["allocation_a"].partner_id.display_name, band, band)
        self.assertIn("Facturas", customers.cell(CIRCUIT_BAND_ROW, 1).value)
        self.assertEqual(customers.cell(DOCUMENT_FIRST_ROW, 1).value, "FA-0001")
        self.assertEqual(
            customers.cell(TRANSFER_FIRST_ROW, 1).value,
            self.scenario["voucher_a"].name,
        )

    def test_sources_have_no_sudo(self):
        """RED/GREEN: the report path MUST NOT escalate privileges."""
        from pathlib import Path

        root = Path(__file__).resolve().parents[1]
        for relative in (
            "report/collection_receipt_assignment_xlsx.py",
            "models/collection_assignment_report.py",
            "models/collection_voucher.py",
        ):
            source = (root / relative).read_text(encoding="utf-8")
            self.assertNotIn("sudo(", source, relative)


class TestReceiptAssignmentXlsxText(TransactionCase):
    """Free text from the board and the ERP must never become a live formula."""

    def test_free_text_is_written_as_text_never_as_formula(self):
        """Names, notes, numbers and banks starting with ``=`` stay strings."""
        injected = {
            "name": '=HYPERLINK("http://evil.example/x","click")',
            "note": "=1+1",
            "number": "=2+2",
            "reference": "=3+3",
            "bank": "=4+4",
            "url": "http://evil.example/account",
        }
        payload = _fixture_payload()
        for row in payload["summary_rows"]:
            row["name"] = injected["name"]
            row["note"] = injected["note"]
        circuit = payload["customer_blocks"][0]["circuits"][0]
        circuit["document_rows"][0]["number"] = injected["number"]
        transfer, check = circuit["voucher_blocks"]
        transfer["rows"][0]["reference"] = injected["reference"]
        transfer["rows"][0]["from_account"] = injected["url"]
        check["rows"][0]["check_bank"] = injected["bank"]
        workbook = load_workbook(BytesIO(builder._build_xlsx(payload)))
        cells = [
            cell
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
        ]
        for label, text in injected.items():
            with self.subTest(field=label):
                matches = [cell for cell in cells if cell.value == text]
                self.assertTrue(matches, f"{text!r} must be written")
                for cell in matches:
                    self.assertEqual(cell.data_type, "s", cell.coordinate)
                    self.assertIsNone(cell.hyperlink, cell.coordinate)
        formulas = [str(cell.value) for cell in cells if cell.data_type == "f"]
        self.assertTrue(formulas, "intentional formulas must still be formulas")
        self.assertFalse([value for value in formulas if "HYPERLINK" in value])
